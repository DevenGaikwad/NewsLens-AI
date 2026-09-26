"""Rebuild the four retained DOCX reports from current Phase 5S Markdown evidence."""

from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
AUTHOR = "Deven Sachin Gaikwad"
BROWN = RGBColor(109, 89, 71)
GREEN = RGBColor(73, 100, 84)
INK = RGBColor(26, 25, 23)
GREY = RGBColor(119, 115, 108)

OUTPUTS = {
    "NewsLens_AI_Project_Report.docx": (
        "NewsLens AI Project Report",
        "Synthetic-only classifier, calibration, evaluation and deployment evidence",
        [
            "README.md",
            "reports/NewsLens_AI_Deployment_and_Audit_Report.md",
            "docs/EDITORIAL_AI_CASE_STUDY.md",
            "docs/LICENSING_STATUS.md",
        ],
        ["reports/figures/class_distribution.png", "reports/figures/confusion_matrix.png", "reports/figures/calibration_reliability.png"],
    ),
    "NewsLens_AI_Code_Explanation_and_Developer_Guide.docx": (
        "NewsLens AI Code Explanation and Developer Guide",
        "Architecture, testing, security and release controls",
        ["docs/ARCHITECTURE.md", "docs/TESTING.md", "docs/PUBLIC_RELEASE_AUDIT.md", "reports/FILE_CHANGE_MANIFEST.md"],
        ["reports/figures/model_comparison.png"],
    ),
    "NewsLens_AI_Complete_Concepts_Methodologies_and_Terminology_Guide.docx": (
        "NewsLens AI Concepts, Methodologies and Terminology Guide",
        "Synthetic data, linear NLP, calibration, leakage controls and responsible interpretation",
        ["docs/DATASET_CARD.md", "docs/MODEL_CARD.md", "docs/PRIVACY.md", "docs/PLACEMENT_INTERVIEW_GUIDE.md"],
        ["reports/figures/feature_importance.png", "reports/figures/roc_pr_curves.png"],
    ),
    "NewsLens_AI_Setup_and_Run_Guide.docx": (
        "NewsLens AI Setup and Run Guide",
        "Python 3.12 local setup, validation and Streamlit deployment",
        ["README.md", "docs/DEPLOYMENT.md", "docs/DEPLOYMENT_CHECKPOINT.md", "docs/DEPLOYMENT_VALUES_TO_FILL.md"],
        [],
    ),
}


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def configure(document: Document, title: str) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.72)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 28, BROWN), ("Heading 1", 18, BROWN), ("Heading 2", 14, GREEN), ("Heading 3", 11, BROWN)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    header = section.header.paragraphs[0]
    header.text = f"NEWSLENS AI  |  {title.upper()}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Calibri"; run.font.size = Pt(8); run.font.bold = True; run.font.color.rgb = GREY
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("© 2026 Deven Sachin Gaikwad · Phase 5S · Page ")
    run.font.name = "Calibri"; run.font.size = Pt(8); run.font.color.rgb = GREY
    page_number(footer)
    document.core_properties.title = title
    document.core_properties.author = AUTHOR
    document.core_properties.subject = "NewsLens AI synthetic-only public model and deployment evidence"


def add_cover(document: Document, title: str, subtitle: str) -> None:
    document.add_paragraph("NEWSLENS AI", style="Subtitle").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("")
    heading = document.add_paragraph(title, style="Title")
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = document.add_paragraph(subtitle)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.size = Pt(14); sub.runs[0].font.color.rgb = GREEN
    document.add_paragraph("")
    note = document.add_paragraph(
        "All classifier training articles and entities are synthetic. This report documents a bounded ledger-consistency demonstration, not unrestricted real-world fact verification."
    )
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in note.runs:
        run.font.italic = True; run.font.color.rgb = GREY
    document.add_paragraph("")
    author = document.add_paragraph(f"Prepared by\n{AUTHOR}\n26 September 2026")
    author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()


def add_inline(paragraph, value: str) -> None:
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", value)
    for part in parts:
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1]); run.font.name = "Consolas"; run.font.size = Pt(8.5)
        elif part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part)


def add_table(document: Document, rows: list[list[str]]) -> None:
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            cell.text = values[column_index].strip() if column_index < len(values) else ""
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Calibri"; run.font.size = Pt(8.5)
                    if row_index == 0: run.font.bold = True
            if row_index == 0: shade(cell, "EAE4D8")
    document.add_paragraph("")


def add_markdown(document: Document, path: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    table_rows: list[list[str]] = []
    in_code = False
    for line in lines:
        if line.startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.25)
            shade_paragraph = OxmlElement("w:shd")
            shade_paragraph.set(qn("w:fill"), "F0EEE9")
            paragraph._p.get_or_add_pPr().append(shade_paragraph)
            run = paragraph.add_run(line or " "); run.font.name = "Consolas"; run.font.size = Pt(8)
            continue
        if line.startswith("|") and line.endswith("|"):
            cells = [cell.strip() for cell in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                continue
            table_rows.append(cells)
            continue
        if table_rows:
            add_table(document, table_rows); table_rows = []
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            document.add_paragraph(stripped[2:], style="Heading 1")
        elif stripped.startswith("## "):
            document.add_paragraph(stripped[3:], style="Heading 2")
        elif stripped.startswith("### "):
            document.add_paragraph(stripped[4:], style="Heading 3")
        elif re.match(r"^\d+\. ", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
        elif stripped.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped)
    if table_rows:
        add_table(document, table_rows)


def build(filename: str, title: str, subtitle: str, sources: list[str], figures: list[str]) -> None:
    document = Document()
    configure(document, title)
    add_cover(document, title, subtitle)
    document.add_heading("Document control", level=1)
    add_table(document, [["Field", "Value"], ["Project", "NewsLens AI"], ["Phase", "5S — synthetic dataset, model and deployment"], ["Author", AUTHOR], ["Status", "PR B evidence package; external deployment values recorded after protected merge"]])
    for index, source in enumerate(sources):
        if index:
            document.add_section(WD_SECTION.NEW_PAGE)
        add_markdown(document, ROOT / source)
        if index == 0 and figures:
            document.add_heading("Selected measured figures", level=2)
            for figure in figures:
                path = ROOT / figure
                if path.exists():
                    document.add_picture(str(path), width=Inches(6.2))
                    caption = document.add_paragraph(path.stem.replace("_", " ").title())
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    caption.runs[0].font.italic = True; caption.runs[0].font.size = Pt(8)
    document.save(DOCS / filename)


def main() -> None:
    for filename, (title, subtitle, sources, figures) in OUTPUTS.items():
        build(filename, title, subtitle, sources, figures)
        print(DOCS / filename)


if __name__ == "__main__":
    main()
